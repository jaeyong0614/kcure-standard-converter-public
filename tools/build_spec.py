"""항목정의서(v2.4) + 보조 워크북(V1.2)에서 변환기 내장 명세(spec/kcure_spec.json)를 추출한다.

사용법:
    python tools/build_spec.py [--final v2.4.xlsx] [--aux V1.2.xlsx]

- 매핑(경로·판정·KR CDI 배정)은 최종본(v2.4)을 기준으로 한다.
- 원천 별칭(암종별 ASIS 항목명), KR CDI 대응유형, 프로파일 canonical, 코드표는 V1.2에서 가져온다.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
import warnings

from openpyxl import load_workbook

warnings.filterwarnings("ignore")

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
ROOT = os.path.dirname(HERE)
from common import setup_console, source_path  # noqa: E402

KRCORE_BASE = "http://www.hl7korea.or.kr/fhir/krcore/StructureDefinition/"
# KR Core V2 상세설명서 8장 예시에서 확인한 프로파일 ID. (확인=원문 확인, 추정=명명규칙 추정)
KRCORE_IDS = {
    "KR Core Patient Profile": ("krcore-patient", "확인"),
    "KR Core Encounter Profile": ("krcore-encounter", "확인"),
    "KR Core Healthcare Organization Profile": ("krcore-healthcare-organization", "확인"),
    "KR Core Condition Profile for Encounter Diagnosis": ("krcore-condition-encounter-diagnosis", "확인"),
    "KR Core Condition Profile for Chief Complaint": ("krcore-condition-chief-complaint", "확인"),
    "KR Core Procedure Profile": ("krcore-procedure", "확인"),
    "KR Core MedicationRequest Profile": ("krcore-medicationrequest", "확인"),
    "KR Core Medication Profile": ("krcore-medication", "확인"),
    "KR Core Observation Profile for Laboratory Result": ("krcore-observation-laboratory-result", "확인"),
    "KR Core DiagnosticReport Profile for Laboratory Results": ("krcore-diagnosticreport-laboratory-results", "확인"),
    "KR Core DiagnosticReport Profile for Diagnostic Imaging": ("krcore-diagnosticreport-diagnostic-imaging", "추정"),
    "KR Core DiagnosticReport Profile for Pathology Results": ("krcore-diagnosticreport-pathology-results", "추정"),
    "KR Core Specimen Profile": ("krcore-specimen", "확인"),
    "KR Core Immunization Profile": ("krcore-immunization", "확인"),
    "KR Core Observation Profile for Vital Signs: Body Weight": ("krcore-bodyweight", "확인"),
    "KR Core Observation Profile for Vital Signs: Body Height": ("krcore-bodyheight", "확인"),
    "KR Core Observation Profile for Vital Signs: Blood Pressure": ("krcore-bloodpressure", "확인"),
}


def norm_key(s):
    """별칭 대조용 정규화: 소문자, 공백·기호 제거."""
    s = str(s or "").lower()
    return re.sub(r"[\s_\-\.\(\)\[\]/·,:;'\"`]+", "", s)


def clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("-", "None") else s


def sheet_rows(wb, name, header_row):
    ws = wb[name]
    rows = list(ws.iter_rows(min_row=header_row, values_only=True))
    head = [clean(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r and r[0] not in (None, ""):
            out.append({h: r[i] if i < len(r) else None for i, h in enumerate(head) if h})
    return out


CODE_PAIR = re.compile(r"^\s*([^:,]+?)\s*:\s*([^,]+?)\s*$")


def parse_allowed(raw, dtype):
    """허용값 문자열 → {codes:[[code,label]], range:[lo,hi], note}"""
    raw = clean(raw)
    res = {}
    if not raw:
        return res
    parts = [p for p in re.split(r"\s*,\s*", raw) if p]
    pairs = [CODE_PAIR.match(p) for p in parts]
    if len(parts) >= 2 and all(pairs):
        res["codes"] = [[m.group(1).strip(), m.group(2).strip()] for m in pairs]
        return res
    m = re.match(r"^\s*(-?\d+(?:\.\d+)?)\s*[~\-]\s*(-?\d+(?:\.\d+)?)\s*$", raw)
    if m and dtype in ("INTEGER", "NUMERIC"):
        res["range"] = [float(m.group(1)), float(m.group(2))]
        return res
    if len(parts) >= 3 and not any(":" in p for p in parts) and dtype == "CODE":
        res["codes"] = [[p, p] for p in parts]
        return res
    return res


def as_int(v):
    try:
        return int(float(str(v)))
    except (TypeError, ValueError):
        return None


def main():
    setup_console()
    ap = argparse.ArgumentParser()
    ap.add_argument("--final", help="항목정의서 최종본 (기본: spec/sources)")
    ap.add_argument("--aux", help="보조 워크북 V1.2 (기본: spec/sources)")
    ap.add_argument("--argo", help="ICGC ARGO 데이터 사전 (기본: spec/sources)")
    ap.add_argument("--mcode", help="mCODE STU4 데이터 사전 (기본: spec/sources)")
    ap.add_argument("--out", default=os.path.join(ROOT, "spec", "kcure_spec.json"))
    a = ap.parse_args()
    a.final = a.final or source_path("final")
    a.aux = a.aux or source_path("aux")
    a.argo = a.argo or source_path("argo")
    a.mcode = a.mcode or source_path("mcode")

    wf = load_workbook(a.final, read_only=True, data_only=True)
    wa = load_workbook(a.aux, read_only=True, data_only=True)

    final = sheet_rows(wf, "10_항목정의_최종", 4)
    master = {r["K-CURE_표준항목ID"]: r for r in sheet_rows(wa, "10_KCURE_항목마스터", 4)}
    map11 = {r["K-CURE_표준항목ID"]: r for r in sheet_rows(wa, "11_표준매핑대장", 4)}

    # 원천(ASIS) 행: Excel 행번호 → 레코드
    asis_ws = wa["01_ASIS_원본"]
    asis = {}
    for i, r in enumerate(asis_ws.iter_rows(min_row=2, values_only=True), start=2):
        asis[i] = {"cancer": clean(r[0]), "name": clean(r[7]), "nameTobe": clean(r[8]),
                   "def": clean(r[9]), "example": clean(r[11])}

    # KR Core canonical 사전 (V1.2 공식명 → URL) + 원문 확인 사전
    core_url = {}
    for r in map11.values():
        n, u = clean(r.get("KR_Core_Profile_공식명")), clean(r.get("KR_Core_Canonical_URL"))
        if n and u.startswith("http"):
            core_url.setdefault(n, u)

    profiles = []
    for r in sheet_rows(wa, "12_FHIR_프로파일후보", 4):
        if not str(r.get("후보ID", "")).startswith("PRF-"):
            continue
        profiles.append({
            "id": r["후보ID"], "name": clean(r.get("K-CURE 프로파일 후보")),
            "resource": clean(r.get("FHIR Resource")), "parent": clean(r.get("KR Core 상위 프로파일")),
            "mcode": clean(r.get("mCODE 참조")), "purpose": clean(r.get("목적")),
            "url": clean(r.get("Canonical_URL_후보")), "domain": clean(r.get("6대도메인")),
        })
    prf_by_id = {p["id"]: p for p in profiles}

    items = []
    cancers = set()
    for r in final:
        iid = clean(r["항목ID"])
        m = master.get(iid, {})
        m11 = map11.get(iid, {})
        dtype = clean(r["데이터타입"]) or "VARCHAR"
        if dtype.startswith("["):
            dtype = "VARCHAR"
        canc = [c.strip() for c in clean(r["출현암종"]).split(",") if c.strip()]
        cancers.update(canc)

        # 별칭: 표준명, 연계항목명, 원천 ASIS 항목명(암종별)
        aliases = {}
        def add_alias(name, cancer=""):
            name = clean(name)
            if not name:
                return
            aliases.setdefault(name, set())
            if cancer:
                aliases[name].add(cancer)
        add_alias(r["표준항목명"])
        add_alias(m.get("연계항목명"))
        add_alias(m.get("TOBE_표준항목명"))
        examples = []
        for rn in str(m.get("원천행번호") or "").split(","):
            rn = as_int(rn)
            if rn and rn in asis:
                a_ = asis[rn]
                add_alias(a_["name"], a_["cancer"])
                add_alias(a_["nameTobe"], a_["cancer"])
                if a_["example"] and len(examples) < 3 and a_["example"] not in examples:
                    examples.append(a_["example"])

        core_prof = clean(r["KR_Core_Profile"])
        core_id = KRCORE_IDS.get(core_prof)
        prf = clean(r["KR-mCODE_프로파일ID"])
        if core_id:
            prof_url, prof_src = KRCORE_BASE + core_id[0], "KR Core(" + core_id[1] + ")"
        elif core_prof.startswith("KR Core 미정의") and prf in prf_by_id:
            prof_url, prof_src = prf_by_id[prf]["url"], "KR-mCODE 후보(" + prf + ")"
        elif core_prof in core_url:
            prof_url, prof_src = core_url[core_prof], "KR Core(V1.2 대장)"
        elif prf in prf_by_id:
            prof_url, prof_src = prf_by_id[prf]["url"], "KR-mCODE 후보(" + prf + ")"
        else:
            prof_url, prof_src = "", ""

        item = {
            "id": iid,
            "name": clean(r["표준항목명"]),
            "def": clean(r["표준정의"]),
            "cat1": clean(r["대분류"]),
            "cat2": clean(r["중분류"]).replace("\n", " "),
            "tier": clean(r["우선검토구분"]),
            "domain": clean(r["mCODE Category"]),
            "role": clean(r["의미역할"]),
            "cancers": canc,
            "type": dtype,
            "len": as_int(r["데이터길이"]),
            "dec": as_int(r["소수점"]),
            "storage": clean(r["저장형식"]),
            "unit": "" if clean(r["단위"]).startswith("[") else clean(r["단위"]),
            "codeYN": clean(r["코드여부"])[:1] if clean(r["코드여부"])[:1] in "YN" else "",
            "codeName": clean(r["표준코드명"]),
            "allowed": clean(r["허용값"]),
            "req": clean(r["필수여부"]),
            "cdi": {
                "cls": "" if clean(r["KR_CDI_Class"]) == "해당 없음" else clean(r["KR_CDI_Class"]),
                "el": "" if clean(r["KR_CDI_Element"]) == "해당 없음" else clean(r["KR_CDI_Element"]),
                "req": "" if clean(r["KR_CDI_필수"]) == "해당 없음" else clean(r["KR_CDI_필수"])[:2].strip(),
                "match": clean(m11.get("KRCDI_대응유형")),
                "no": as_int(m11.get("KRCDI_항목No")),
            },
            "core": {"profile": core_prof, "path": clean(r["KR_Core_FHIR_Path"]),
                     "judge": clean(r["KR_Core_판정"]), "url": prof_url, "urlSrc": prof_src},
            "mcode": {"profile": clean(r["mCODE_Profile"]), "path": clean(r["mCODE_FHIR_Path"]),
                      "level": clean(r["mCODE_참조수준"])},
            "prf": prf,
            "grade": clean(r["심사등급"]),
            "stage": clean(r["매핑확정단계"]),
            "gap": clean(r["정의결손코드"]),
            "examples": examples,
            "aliases": [[k, sorted(v)] for k, v in aliases.items()],
        }
        if not item["cdi"]["el"]:
            item["cdi"] = {}
        item.update(parse_allowed(r["허용값"], dtype))
        items.append(item)

    krcdi = []
    m21 = {as_int(r["항목No"]): r for r in sheet_rows(wf, "21_KRCDI_V3_항목마스터", 4)}
    for r in sheet_rows(wf, "20_KRCDI_KRCore_대응표", 4):
        no = as_int(r["항목No"])
        d = m21.get(no, {})
        krcdi.append({
            "no": no, "cls": clean(r["분류(한글)"]), "clsEn": clean(r["분류(영문)"]),
            "el": clean(r["항목(Element)"]), "elEn": clean(d.get("영문명")), "def": clean(d.get("개념 정의")),
            "req": clean(r["필수여부"])[:2].strip(), "type": clean(r["데이터타입"]),
            "term": clean(r["참조 용어체계(원문)"]), "coreProfile": clean(r["1순위 KR Core Profile"]),
            "corePath": clean(r["1순위 FHIR Path"]), "alt": clean(r["허용 대안"]), "cond": clean(r["선택 조건"]),
        })

    # ICGC ARGO 후보 매핑 (항목정의서에 없는 대응 → 규칙 기반 자동 후보)
    import argo_map
    argo_dict, n_argo = argo_map.apply(items, a.argo)
    # mCODE STU4 데이터 사전: 프로파일 URL·데이터요소 연결 (대응 자체는 항목정의서 v2.4)
    import mcode_dict
    mcode_d, n_mp, n_me = mcode_dict.apply(items, a.mcode)
    # OMOP CDM v5.4 후보 매핑 (항목정의서에 없는 대응 → 규칙 기반 자동 후보)
    import omop_map
    omop_d = omop_map.apply(items)

    code_systems = {}
    for r in sheet_rows(wa, "05_코드값_정의", 1):
        code_systems.setdefault(clean(r["표준코드명"]), []).append(
            [clean(r["코드값"]), clean(r["코드값명"]), clean(r["코드설명"])])

    spec = {
        "meta": {
            "title": "K-CURE 9개 암종 항목 → KR CDI·KR Core·DCAT 변환 명세",
            "source": os.path.basename(a.final),
            "aux": os.path.basename(a.aux),
            "generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
            "itemCount": len(items),
            "cancers": sorted(cancers),
            "stageNote": "규칙승인(후보) 단계 — 항목별 전문가 의미검증(L3) 0건",
            "krcoreBase": KRCORE_BASE,
            "krcoreIds": {k: list(v) for k, v in KRCORE_IDS.items()},
            "kcureBase": "http://khis.kr/fhir/kcure/",
        },
        "items": items,
        "krcdi": krcdi,
        "profiles": profiles,
        "codeSystems": code_systems,
        "argo": argo_dict,
        "mcode": mcode_d,
        "omop": omop_d,
    }
    spec["meta"]["argoNote"] = f"ICGC ARGO v{argo_dict['version']} 매핑은 규칙 기반 자동 후보({n_argo}개 항목) — 전문가 검토 전"
    spec["meta"]["omopNote"] = "OMOP CDM v5.4 매핑은 규칙 기반 자동 후보(1,255개 항목, 테이블·필드 수준) — 개념 ID(concept_id) 매핑은 Athena 어휘로 별도 수행"
    spec["meta"]["mcodeNote"] = f"mCODE {mcode_d['version']} 대응은 항목정의서 v2.4 기준(규칙승인 후보), 프로파일 URL·데이터요소는 mCODE 데이터 사전 연결({n_mp}개 항목 프로파일, {n_me}개 항목 데이터요소)"
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    with open(a.out, "w", encoding="utf-8") as f:
        json.dump(spec, f, ensure_ascii=False, separators=(",", ":"))
    # file:// 에서도 <script src>로 읽을 수 있도록 JS 래퍼를 함께 쓴다.
    with open(os.path.splitext(a.out)[0] + ".js", "w", encoding="utf-8") as f:
        f.write("window.KCURE_SPEC=")
        json.dump(spec, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")
    n_alias = sum(len(i["aliases"]) for i in items)
    print(f"items={len(items)} aliases={n_alias} krcdi={len(krcdi)} profiles={len(profiles)} "
          f"codeSystems={len(code_systems)} size={os.path.getsize(a.out)//1024}KB -> {a.out}")


if __name__ == "__main__":
    main()
